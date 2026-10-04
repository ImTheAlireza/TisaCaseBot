"""
processor.py — هسته‌ی پردازش فایل سفارش برای ربات تلگرام

ورودی: فایل .xlsx / .csv / .pdf  (جدول سفارش‌های تیساکیس / تیسا چاپ)
       ستون‌های جدول: ردیف | بارکد | تاریخ ثبت | نام گیرنده | کد سفارش | مقصد | نام فروشگاه | آدرس | وزن
       (فایل‌هایی با هدر انگلیسی order_id / tracking_code هم پشتیبانی می‌شوند)

خروجی: یک ``Report`` (دیتاکلس) با:
  - csv_text      : فایل tracking.csv با ستون‌های order_id,tracking_code
                    (فقط ردیف‌هایی که بارکدشان معتبر است)
  - summary       : خلاصه‌ی گزارش برای نمایش در چت
  - problems_csv  : گزارش کامل مشکلات، با لینک سطرِ مبدأ (یا None اگر مشکلی نبود)
  - review_csv/review_xlsx : ردیف‌هایی که باید دیده شوند (خطا و هشدار); نسخهٔ xlsx طوری است که ستون بارکد
                    در اکسل «متن» می‌ماند و عدد ۲۴ رقمی خراب نمی‌شود — و چون هدرهایش
                    همان «بارکد / کد سفارش / ردیف» است، خودِ ربات هم می‌تواند فایلِ
                    اصلاح‌شده را بخواند (رفت‌وبرگشت، بدون مسیرِ دوم).
  - questions     : اگر دو ستون محتمل بود، سؤالِ «کدام ستون؟» — حدس زدن ممنوع.

هیچ‌چیز اینجا به تلگرام وابسته نیست؛ تست‌هایش هم بدون ربات اجرا می‌شوند.
"""

from __future__ import annotations

import codecs
import io
import logging
import os
import re
from dataclasses import dataclass, field, fields
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

import pandas as pd

from bot.config import settings
from bot.services import barcodes

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# نرمال‌سازی متن (ارقام فارسی/عربی → انگلیسی، حذف فاصله/نیم‌فاصله/کاما)
# ---------------------------------------------------------------------------
_FA2EN = {
    "۰": "0",
    "۱": "1",
    "۲": "2",
    "۳": "3",
    "۴": "4",
    "۵": "5",
    "۶": "6",
    "۷": "7",
    "۸": "8",
    "۹": "9",
    "٠": "0",
    "١": "1",
    "٢": "2",
    "٣": "3",
    "٤": "4",
    "٥": "5",
    "٦": "6",
    "٧": "7",
    "٨": "8",
    "٩": "9",
    "ي": "ی",
    "ك": "ک",
    "ة": "ه",
    "ۀ": "ه",
}
# str.translate نیاز به جدول ordinal دارد — دیکشنری str→str بی‌اثر است
_FA2EN_TABLE = str.maketrans(_FA2EN)


def _norm(s) -> str:
    """نرمال‌سازی برای مقایسه‌ی نام ستون‌ها (هدرها)"""
    if s is None:
        return ""
    if isinstance(s, float) and s.is_integer():
        s = int(s)
    s = str(s).translate(_FA2EN_TABLE)
    s = re.sub(r"[\s\u00a0\u200c\u200f\u202a\u202b,]", "", s)
    return s.lower()


def _clean(v) -> str:
    """نرمال‌سازی مقدار یک سلول (بدون lower کردن — برای داده)"""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, float) and pd.isna(v):
        return ""
    s = str(v)
    s = re.sub(r"[\s\u00a0\u200c\u200f\u202a\u202b,]", "", s)
    return s.translate(_FA2EN_TABLE)


# ---------------------------------------------------------------------------
# نام ستون‌ها (بعد از نرمال‌سازی — فاصله‌ها حذف شده‌اند)
# ---------------------------------------------------------------------------
#: What a column has to be called to be *the* barcode column. «کد رهگیری» is in here
#: because that is what the other shop's exports call the same 24-digit thing — and a
#: sheet that has both columns is exactly the case where the bot must ask, not pick.
_BARCODE_HEADERS = {
    "بارکد",
    "باركد",
    "barcode",
    "trackingcode",
    "tracking_code",
    "tracking",
    "کدرهگیری",
}
_CODE_HEADERS = {"کدسفارش", "کدسفارشگیرنده", "orderid", "order_id", "order", "کد"}
_ROW_HEADERS = {"ردیف", "رديف", "ردی", "شماره", "no"}
# ستون نام گیرنده — وقتی ستون «کد سفارش» جدا وجود ندارد، کد داخل همین ستون است
# (مثل «امیرحسین عاشوری ۳۰۶۱۷۶»)
_NAME_HEADERS = {"نامگ", "نامگیرنده", "گیرنده", "نامونامخانوادگیگیرنده", "recipient"}

FIELD_BARCODE = "barcode"
FIELD_CODE = "code"

# Which attribute of ``Layout`` a field lives in: only this module has to know it, so the
# flow can speak in fields and never name a column attribute of its own.
_LAYOUT_ATTR = {FIELD_BARCODE: "barcode_col", FIELD_CODE: "code_col"}

_RE_DATE = re.compile(r"14\d{2}/\d{2}/\d{2}")
# کد ۵-۶ رقمی داخل متن (مثلاً چسبیده به نام گیرنده)
_RE_CODE_IN_TEXT = re.compile(r"(?<![0-9])([0-9]{5,6})(?![0-9])")

#: How many columns a question may offer — a Telegram keyboard with 30 buttons is
#: not a question, it is a chore.
MAX_OPTIONS = 8


class RowLimitError(ValueError):
    """More rows than ``MAX_ROWS`` — refuse instead of grinding the bot to dust."""


class MemoryLimitError(RowLimitError):
    """The file could not be read because the worker ran out of memory.

    A ``RowLimitError`` on purpose: the flow already answers that family with «🚧 …» and
    something to *do*, which is exactly the right channel. The sentence itself names both
    causes — a file that is genuinely too big, and a host whose memory ceiling is too low
    for even a small file — because from inside the parser the two look identical, and only
    the shop can tell which one it is (the worker log prints the numbers).
    """


# ---------------------------------------------------------------------------
# What the file looks like (schema detection)
# ---------------------------------------------------------------------------
@dataclass(frozen=True)
class Column:
    """One candidate column, as the user has to recognise it in a button."""

    index: int
    header: str
    sample: str

    def label(self) -> str:
        """``C · «بارکد» · 610001573845…`` — short enough for a button."""
        head = self.header or "(بدون هدر)"
        if len(head) > 18:
            head = head[:17] + "…"
        sample = self.sample[:10] + ("…" if len(self.sample) > 10 else "")
        tail = f" · {sample}" if sample else ""
        return f"{_col_letter(self.index)} · «{head}»{tail}"


@dataclass(frozen=True)
class Question:
    """«کدام ستون؟» — asked instead of guessing when two headers are plausible."""

    field: str  # FIELD_BARCODE | FIELD_CODE
    prompt: str
    options: tuple[Column, ...]
    allow_skip: bool = False

    def answers(self, index: int) -> dict[str, Any]:
        return {self.field: index}


@dataclass(frozen=True)
class Layout:
    """The columns this file will be read through, and how each was decided."""

    header_row: int = -1  # 0-based row index of the header, -1 = none
    barcode_col: int | None = None
    code_col: int | None = None
    row_col: int | None = None
    name_col: int | None = None
    sheet: str = ""
    #: «هدر» when the header said so, «انتخاب تو» when the user picked it
    decided: tuple[tuple[str, str], ...] = field(default_factory=tuple)

    def with_choice(self, field: str, index: int) -> Layout:
        """This layout with one column replaced by what the user picked.

        ``-1`` means «بدون این ستون» — the field is switched off, not moved. The label is
        recorded with it, so a summary can still tell a header match from a human choice.
        """
        attrs = {f.name: getattr(self, f.name) for f in fields(self) if f.name != "decided"}
        attrs[_LAYOUT_ATTR[field]] = None if index < 0 else index
        labels = dict(self.decided)
        labels[field] = "انتخاب تو"
        return Layout(**attrs, decided=tuple(sorted(labels.items())))

    def as_dict(self) -> dict[str, Any]:
        return {
            "header_row": self.header_row,
            "barcode_col": self.barcode_col,
            "code_col": self.code_col,
            "row_col": self.row_col,
            "name_col": self.name_col,
            "sheet": self.sheet,
            "decided": list(self.decided),
        }

    @classmethod
    def from_dict(cls, data: dict[str, Any]) -> Layout:
        return cls(
            header_row=int(data.get("header_row", -1)),
            barcode_col=_as_int(data.get("barcode_col")),
            code_col=_as_int(data.get("code_col")),
            row_col=_as_int(data.get("row_col")),
            name_col=_as_int(data.get("name_col")),
            sheet=str(data.get("sheet") or ""),
            decided=tuple(tuple(x) for x in (data.get("decided") or ())),
        )

    def schema_line(self) -> str:
        """Tell the user which columns were used — never read a file silently.

        A column the owner switched off is reported too: «کد سفارش ← ندارد (انتخاب تو)»
        is the difference between «the bot forgot my code column» and «I told it not to
        read one».
        """
        decided = dict(self.decided)
        parts = []
        for name, label in (("barcode", "بارکد"), ("code", "کد سفارش"), ("row", "ردیف")):
            how = decided.get(name)
            if how is None:
                continue
            index = getattr(self, f"{name}_col" if name != "row" else "row_col")
            parts.append(
                f"{label} ← {_col_letter(index)} ({how})" if index is not None else f"{label} ← ندارد ({how})"
            )
        return " · ".join(parts)


def _as_int(value: object) -> int | None:
    try:
        return None if value is None else int(str(value))
    except (TypeError, ValueError):
        return None


def _col_letter(index: int) -> str:
    """0-based column index → Excel letter (0→A, 25→Z, 26→AA)."""
    letters = ""
    n = int(index) + 1
    while n:
        n, rem = divmod(n - 1, 26)
        letters = chr(ord("A") + rem) + letters
    return letters or "A"


# ---------------------------------------------------------------------------
# خواندن فایل — همان چیزی که از سامانه می‌آید، نه ایده‌آلِ ما
# ---------------------------------------------------------------------------
class UnreadableFileError(ValueError):
    """فایل خوانده نشد — با دلیلی که کاربر بتواند کاری کند.

    Every parser-level failure (a truncated zip, a web panel's HTML named ``.xlsx``, a
    scanned PDF) is turned into one Persian sentence here; the English/technical text
    stays in the log, where a developer can find it.
    """


#: Delimiters a Persian Windows export can use. Excel in a Persian locale writes «;» —
#: reading that as one column is how a full order file became «ستون بارکد پیدا نشد».
_DELIMITERS = (",", ";", "\t", "|")
#: Shorter than this cannot be a barcode: Tisa's are 24 digits, GTINs 8/12/13/14.
_MIN_BARCODE_DIGITS = 8
#: How many sheets are previewed before one is chosen (a 300-sheet workbook must not
#: turn one upload into 300 reads).
_SHEET_PREVIEWS = 12
#: Lines of a printed table share a baseline within this many points. Rounding to 0.1pt
#: (the previous version) split one row into several «lines» whenever a cell's font
#: differed, and the row disappeared from the output.
_LINE_TOLERANCE = 3.0


def _decode_bytes(raw: bytes) -> str:
    """متنِ فایل با هر کدگذاری‌ای که خروجی‌های واقعی دارند.

    cp1256 is what a Persian Windows tool writes; utf-16 (with or without a BOM) is what
    «Unicode Text» and several web panels write; utf-8 is tried first because a cp1256
    byte string that happens to be valid utf-8 is far rarer than the reverse.
    """
    try:
        if raw.startswith(codecs.BOM_UTF8):
            return raw.decode("utf-8-sig")
        if raw.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
            return raw.decode("utf-16")
        if raw.find(b"\x00") >= 0:  # UTF-16 without a BOM (ASCII text has NUL bytes)
            for encoding in ("utf-16-le", "utf-16-be"):
                try:
                    return raw.decode(encoding)
                except UnicodeDecodeError:
                    continue
    except UnicodeDecodeError:  # a BOM that lies about the bytes behind it
        pass
    for encoding in ("utf-8", "cp1256", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")  # every byte maps — a last resort that cannot fail


def _head_lines(text: str, limit: int) -> list[str]:
    """The first ``limit`` non-empty lines — without materialising a million-line list.

    A 25 MB export is also a list of a million strings if we let ``splitlines()`` build
    one; the delimiter and the width only need the top of the file.
    """
    lines: list[str] = []
    for line in io.StringIO(text):
        if line.strip():
            lines.append(line)
            if len(lines) >= limit:
                break
    return lines


def _sniff_delimiter(text: str) -> str:
    """جداکنندهٔ واقعی فایل: «,» و «;» و tab و «|».

    Counting beats ``csv.Sniffer`` for this job: the sniffer guesses from one sample and
    gives up on a short sheet, while the delimiter that appears the same number of times
    on most lines is what «CSV» means in practice. Comma wins ties, because it is the
    default in every tool.
    """
    sample = _head_lines(text, 8)
    if not sample:
        return ","
    best, best_score = ",", (0.0, 0)
    for delimiter in _DELIMITERS:
        counts = [len(re.sub(r'"[^"]*"', "", line).split(delimiter)) - 1 for line in sample]
        present_counts = [count for count in counts if count > 0]
        if not present_counts:
            continue
        mode = max(set(present_counts), key=present_counts.count)
        score = (
            len(present_counts) / len(counts) + present_counts.count(mode) / len(counts),
            len(present_counts),
        )
        if score > best_score:
            best, best_score = delimiter, score
    return best


def _rows_to_frame(rows: list[list[str]]) -> pd.DataFrame:
    """A grid of text → DataFrame; ragged rows are padded, because exports are ragged."""
    width = max((len(row) for row in rows), default=0)
    if width == 0:
        return pd.DataFrame()
    return pd.DataFrame([row + [""] * (width - len(row)) for row in rows], dtype=object)


class _HtmlTableParser(HTMLParser):
    """جدول‌های یک فایل HTML را به گرید تبدیل می‌کند (stdlib؛ بدون lxml).

    Web panels in this market export «Excel» that is really an HTML table, and some
    banks and print services hand back the same. pandas cannot read those without
    lxml/html5lib — and the order file must not be refused because of it.
    """

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[str]]] = []
        self._table: list[list[str]] | None = None
        self._row: list[str] | None = None
        self._cell: list[str] | None = None
        self._colspan = 1

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attributes = dict(attrs)
        if tag == "table":
            self._table = []
        elif tag == "tr" and self._table is not None:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []
            try:
                self._colspan = max(1, int(str(attributes.get("colspan") or "1")))
            except ValueError:
                self._colspan = 1
        elif tag == "br" and self._cell is not None:
            self._cell.append(" ")

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag in ("td", "th") and self._cell is not None and self._row is not None:
            # Cell text is collapsed to one line: a wrapped address is still one cell.
            text = " ".join("".join(self._cell).split())
            self._row.extend([text] + [""] * (self._colspan - 1))
            self._cell = None
        elif tag == "tr" and self._row is not None and self._table is not None:
            if any(cell.strip() for cell in self._row):
                self._table.append(self._row)
            self._row = None
        elif tag == "table" and self._table is not None:
            if self._table:
                self.tables.append(self._table)
            self._table = None


def _html_rows(text: str) -> list[list[str]] | None:
    """The biggest table of an HTML document — the order table is the long one."""
    parser = _HtmlTableParser()
    try:
        parser.feed(text)
        parser.close()
    except Exception:  # pragma: no cover — a malformed tag soup, not a file we can use
        return None
    if not parser.tables:
        return None
    return max(parser.tables, key=len)


def _spreadsheetml_rows(text: str) -> list[list[str]] | None:
    """Excel 2003 XML (SpreadsheetML) — «ذخیره به‌صورت XML» در پنل‌های قدیمی."""
    import xml.etree.ElementTree as ET

    try:
        root = ET.fromstring(text)
    except ET.ParseError:
        return None
    namespace = "urn:schemas-microsoft-com:office:spreadsheet"
    best: list[list[str]] = []
    for table in root.iter(f"{{{namespace}}}Table"):
        rows: list[list[str]] = []
        for row in table.findall(f"{{{namespace}}}Row"):
            cells: list[str] = []
            for cell in row.findall(f"{{{namespace}}}Cell"):
                index = cell.get(f"{{{namespace}}}Index")
                if index and index.isdigit() and int(index) > len(cells) + 1:
                    cells.extend([""] * (int(index) - len(cells) - 1))
                data = cell.find(f"{{{namespace}}}Data")
                cells.append("".join(data.itertext()) if data is not None else "")
            if any(cell.strip() for cell in cells):
                rows.append(cells)
        if len(rows) > len(best):
            best = rows
    return best or None


def _unreadable(what: str, exc: Exception) -> UnreadableFileError:
    """یک خطای پارسر → جملهٔ فارسی؛ نوع خطا هم می‌آید تا قابل گزارش باشد.

    «راه‌حل» این‌جا نیست: جریانِ چت آن را با توجه به پسوند فایل اضافه می‌کند تا هر
    پیامِ خطا دقیقاً یک جملهٔ «چه بفرست» داشته باشد.
    """
    return UnreadableFileError(f"{what} (خطای فنی: {type(exc).__name__}).")


def _memory_error(exc: BaseException) -> MemoryLimitError:
    """«MemoryError» خام → جملهٔ صریح با هر دو علتِ ممکن.

    از داخلِ پارسر، «فایل واقعاً بزرگ است» و «سقفِ حافظهٔ کارگر برای همین هاست کم است»
    یک‌شکل دیده می‌شوند؛ پس هر دو گفته می‌شوند و عددها در ``logs/bot.log`` می‌مانند.
    کاربری که فایلِ ۲۲KB فرستاده بود، این‌جا دیگر «خطای فنی: MemoryError» نمی‌بیند.
    """
    return MemoryLimitError(
        "حافظهٔ پردازشِ ربات برای خواندن این فایل پر شد. "
        f"(خطای فنی: {type(exc).__name__})\n"
        "• اگر فایل بزرگ است: به دو یا چند بخش تقسیمش کن، یا همان گزارش را CSV/PDF بگیر.\n"
        "• اگر فایل کوچک است: سقفِ حافظهٔ کارگر کم است — WORKER_MEMORY_MB را در .env "
        "بالا ببر (یا حافظهٔ سرور را). اعدادِ دقیق در logs/bot.log نوشته شده‌اند."
    )


def _pick_sheet(book: pd.ExcelFile, names: list[str]) -> tuple[str, tuple[str, ...]]:
    """برگه‌ای که جدول سفارش‌ها در آن است — نه همیشه برگهٔ اول.

    A multi-sheet export is normal (a «راهنما» sheet first, the orders on the second),
    and reading sheet 1 blindly answered «ستون بارکد پیدا نشد» for a file whose table
    was one tab away. A sheet with the shop's own headers beats one with more rows.
    """
    best_name, best_score = names[0], -1
    for name in names[:_SHEET_PREVIEWS]:
        try:
            preview = pd.read_excel(book, sheet_name=name, header=None, dtype=object, nrows=25)
        except Exception:  # pragma: no cover — one broken sheet must not lose the file
            logger.debug("sheet %r preview failed", name, exc_info=True)
            continue
        header_points, data_rows = 0, 0
        for cells in preview.to_numpy().tolist():
            normed = [_norm(cell) for cell in cells]
            if any(cell in _BARCODE_HEADERS for cell in normed):
                header_points += 2
            if any(cell in _CODE_HEADERS for cell in normed):
                header_points += 1
            if any(str(cell).strip() for cell in cells):
                data_rows += 1
        if not data_rows:
            continue  # an empty sheet is never the answer
        score = header_points * 1000 + min(data_rows, 999)
        if score > best_score:
            best_name, best_score = name, score
    if best_name == names[0]:
        return best_name, ()
    return best_name, (f"جدول در برگهٔ «{best_name}» بود (نه «{names[0]}»).",)


def _read_not_really_xlsx(path, exc: Exception) -> tuple[pd.DataFrame, str, tuple[str, ...]]:
    """A file named ``.xlsx`` that is not a zip: HTML, XML, an old ``xls``, or broken.

    Before this, every one of those reached the warehouse as «File is not a zip file» —
    a sentence that tells nobody what to send instead.
    """
    try:
        raw = Path(path).read_bytes()
        text = _decode_bytes(raw)
        rows = _spreadsheetml_rows(text) if "spreadsheet" in text[:3000].lower() else None
        if rows is None:
            rows = _html_rows(text)
    except MemoryError as oom:  # a web panel's HTML export can be absurdly large
        raise _memory_error(oom) from oom
    if rows is not None:
        logger.info("xlsx is really an HTML/XML table; read as a table: %s", path)
        return (
            _rows_to_frame(rows),
            "",
            ("فایل پسوند اکسل داشت ولی محتوایش جدول HTML/XML بود؛ همان جدول خوانده شد.",),
        )
    if raw[:4] == b"\xd0\xcf\x11\xe0":
        raise UnreadableFileError("این فایل اکسل قدیمی (xls 97-2003) است، نه xlsx/xlsm.")
    raise _unreadable("فایل اکسل سالم نیست — خراب یا نصفه دانلود شده", exc)


def _read_excel(path) -> tuple[pd.DataFrame, str, tuple[str, ...]]:
    """برگهٔ درستِ فایل اکسل، به‌همراه نامش (برای لینک سلول‌ها) و یادداشت‌های خواندن."""
    import zipfile

    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > 2000 or sum(item.file_size for item in members) > 128 * 1024 * 1024:
                raise RowLimitError("حجم بازشده/تعداد فایل‌های XLSX از سقف ایمن بیشتر است")
    except zipfile.BadZipFile as exc:
        return _read_not_really_xlsx(path, exc)
    except MemoryError as exc:
        raise _memory_error(exc) from exc
    try:
        book = pd.ExcelFile(path, engine="openpyxl")
    except MemoryError as exc:
        raise _memory_error(exc) from exc
    except Exception as exc:
        logger.warning("xlsx could not be opened: %s", path, exc_info=True)
        raise _unreadable("این فایل به‌عنوان اکسل باز نشد", exc) from exc
    with book:
        names = [str(name) for name in book.sheet_names]
        if not names:
            raise UnreadableFileError("این فایل اکسل هیچ برگه‌ای ندارد.")
        chosen, notes = _pick_sheet(book, names)
        try:
            frame = pd.read_excel(
                book, sheet_name=chosen, header=None, dtype=object, nrows=settings.max_rows + 1
            )
        except MemoryError as exc:
            raise _memory_error(exc) from exc
        except Exception as exc:  # pragma: no cover — a broken sheet in a readable book
            logger.warning("sheet %r could not be read: %s", chosen, path, exc_info=True)
            raise _unreadable(f"خواندن برگهٔ «{chosen}» ممکن نشد", exc) from exc
    return frame, chosen, notes


def _csv_width(text: str, delimiter: str, lines: int = 200) -> int:
    """How many fields this CSV really has — counted, so short rows are padded, not lost.

    Without an explicit width pandas takes the *first* line as the truth: a report whose
    first line is a title («گزارش سفارش‌ها») then reads every later row as a broken line
    and the whole table disappears.
    """
    width = 1
    for line in _head_lines(text, lines):
        width = max(width, len(re.sub(r'"[^"]*"', "", line).split(delimiter)))
    return min(width, 512)


def _read_text_table(path) -> tuple[pd.DataFrame, str, tuple[str, ...]]:
    """CSV با هر کدگذاری/جداکننده‌ای که خروجی‌های فارسی دارند."""
    try:
        text = _decode_bytes(Path(path).read_bytes())
    except MemoryError as exc:
        raise _memory_error(exc) from exc
    delimiter = _sniff_delimiter(text)
    names = list(range(_csv_width(text, delimiter)))
    notes: list[str] = []
    kwargs: dict[str, Any] = {
        "sep": delimiter, "header": None, "names": names, "dtype": object,
        "keep_default_na": False, "nrows": settings.max_rows + 1,
    }
    try:
        frame = pd.read_csv(io.StringIO(text), **kwargs)
    except MemoryError as exc:
        raise _memory_error(exc) from exc
    except pd.errors.ParserError:
        # A ragged export is common (a note row with two fields). Long lines are the one
        # thing this cannot repair, so they are skipped — and the report says so.
        logger.warning("csv has broken lines, skipping them: %s", path, exc_info=True)
        try:
            frame = pd.read_csv(io.StringIO(text), engine="python", on_bad_lines="skip", **kwargs)
        except MemoryError as inner:
            raise _memory_error(inner) from inner
        except Exception as inner:
            raise _unreadable("فایل CSV خوانده نشد", inner) from inner
        notes.append("چند خط شکستهٔ فایل CSV خوانده نشد؛ بهتر است خروجی را دوباره از سامانه بگیری.")
    except Exception as exc:
        logger.warning("csv could not be read: %s", path, exc_info=True)
        raise _unreadable("فایل CSV خوانده نشد (کدگذاری یا جداکنندهٔ ناشناس)", exc) from exc
    frame = frame.fillna("")  # fields a short row did not have are empty cells, not «nan»
    if delimiter != ",":
        # Tab in a chat bubble reads as nothing at all; name it.
        shown = "tab" if delimiter == "\t" else delimiter
        notes.append(f"جداکنندهٔ فایل «{shown}» بود و همان‌طور خوانده شد.")
    return frame, "", tuple(notes)


def _read_table(path, ext: str) -> tuple[pd.DataFrame, str, tuple[str, ...]]:
    """One entry point for both file kinds: (frame, sheet name, notes)."""
    if ext in (".xlsx", ".xlsm"):
        return _read_excel(path)
    return _read_text_table(path)


def _shape_score(df: pd.DataFrame, col: int, header_row: int) -> tuple[int, str]:
    """How much this column looks like barcodes, and its first non-empty value.

    Only used to *rank the options of a question* — never to pick a column by
    itself. A 24-digit number is the fingerprint of a Tisa tracking code.
    """
    long_digits = 0
    any_digits = 0
    sample = ""
    for r in range(header_row + 1, min(len(df), header_row + 60)):
        v = _clean(_cell(df, r, col))
        if not v:
            continue
        if not sample:
            sample = v
        if v.isdigit():
            any_digits += 1
            if len(v) >= 12:
                long_digits += 1
    return (long_digits * 2 + any_digits), sample


def _headers_of(df: pd.DataFrame, row: int) -> list[str]:
    """The header cells of ``row``; empty when the file has no header row."""
    if row < 0 or row >= len(df):
        return []
    return [str(c).strip() for c in df.iloc[row].tolist()]


def scan_table(
    df: pd.DataFrame,
    answers: dict[str, int | None] | None = None,
    *,
    sheet: str = "",
    prior_labels: dict[str, str] | None = None,
) -> tuple[Layout, tuple[Question, ...]]:
    """Find the header row and the columns, or ask which one is meant.

    ``answers`` carries what the user already picked (``-1`` means «این ستون نیست»
    / «بدون این ستون»); a picked column is never overruled by a header, and an
    ambiguous one is never guessed. ``prior_labels`` says how each inherited column was
    decided the first time, so re-reading a file after an answer does not relabel the
    columns nobody touched.
    """
    inherited = dict(prior_labels or {})
    answers = {k: v for k, v in (answers or {}).items() if v is not None}

    header_rows = []
    for i in range(min(20, len(df))):
        cells = [_norm(c) for c in df.iloc[i].tolist()]
        if any(c in _BARCODE_HEADERS for c in cells) or any(c in _CODE_HEADERS for c in cells):
            header_rows.append(i)
    # -1 when the file has no header at all: then row 1 is data, not a label.
    picked_row = answers.get("header_row")
    header_row = int(picked_row) if picked_row is not None else (header_rows[0] if header_rows else -1)
    raw_headers = _headers_of(df, header_row)
    normed = [_norm(c) for c in raw_headers]

    def candidates(pool: set[str]) -> list[Column]:
        found = [j for j, c in enumerate(normed) if c in pool]
        return [Column(j, raw_headers[j], _clean(_cell(df, header_row + 1, j))) for j in found]

    shapes = {j: _shape_score(df, j, header_row) for j in range(max(len(raw_headers), df.shape[1]))}
    decided: dict[str, str] = {}
    questions: list[Question] = []

    barcode = answers.get(FIELD_BARCODE)
    if barcode is not None:
        decided[FIELD_BARCODE] = inherited.get(FIELD_BARCODE, "انتخاب تو")
    else:
        pool = candidates(_BARCODE_HEADERS)
        if len(pool) == 1:
            barcode = pool[0].index
            decided[FIELD_BARCODE] = "هدر"
        else:
            # Ask about what is plausible, not about every column in the sheet: the
            # two barcode-named ones, or — when nothing is named — the columns whose
            # values look like barcodes.
            options = pool or [
                Column(j, raw_headers[j] if j < len(raw_headers) else "", shapes[j][1])
                for j in sorted(shapes, key=lambda j: -shapes[j][0])[:MAX_OPTIONS]
                if shapes[j][1] or j < len(raw_headers)
            ]
            if not options:
                raise ValueError(
                    "ستون «بارکد» در فایل پیدا نشد. فایل باید جدول سفارش‌ها "
                    "(ردیف / بارکد / تاریخ ثبت / نام گیرنده / کد سفارش / …) باشد."
                )
            why = "دو ستون با نام بارکد وجود دارد" if len(pool) > 1 else "هیچ ستونی «بارکد» نام ندارد"
            questions.append(
                Question(
                    FIELD_BARCODE,
                    f"{why} — کدام ستون بارکد/کد رهگیری است؟",
                    tuple(options[:MAX_OPTIONS]),
                )
            )

    code = answers.get(FIELD_CODE)
    if code is not None:
        decided[FIELD_CODE] = inherited.get(FIELD_CODE, "انتخاب تو")
    else:
        pool = candidates(_CODE_HEADERS)
        if len(pool) == 1:
            code = pool[0].index
            decided[FIELD_CODE] = "هدر"
        elif len(pool) > 1:
            questions.append(
                Question(
                    FIELD_CODE,
                    "چند ستون شبیه «کد سفارش» است — کدام را بردارم؟ "
                    "اگر هیچ‌کدام نبود، «بدون کد سفارش» را بزن.",
                    tuple(pool),
                    allow_skip=True,
                )
            )

    rows_pool = candidates(_ROW_HEADERS)
    name_pool = candidates(_NAME_HEADERS)
    layout = Layout(
        header_row=header_row,
        barcode_col=None if barcode is not None and barcode < 0 else _as_int(barcode),
        code_col=None if code is not None and code < 0 else _as_int(code),
        row_col=rows_pool[0].index if rows_pool else None,
        name_col=name_pool[0].index if name_pool else None,
        sheet=sheet,
        decided=tuple(sorted(decided.items())),
    )
    return layout, tuple(questions)


def _cell(df, row_index, col):
    """یک سلول با احتساب ستونِ ناموجود و مقدار NaN."""
    if col is None:
        return ""
    if col >= df.shape[1] or row_index >= len(df):
        return ""
    v = df.iat[row_index, col]
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    return v


def _collect_rows(
    df, header_idx, bc_col, code_col, row_col, name_col, *, name_for_code: bool = True
) -> list[dict[str, Any]]:
    """سطرهای داده.

    سلول‌ها با ``_cell`` خوانده می‌شوند (نه یک closure داخل حلقه): نسخهٔ قبلی
    متغیر حلقه را دیر-بند می‌کرد که با هر تغییر کوچک، مقدار سطرِ اشتباه می‌دهد.

    ``name_for_code`` بستنِ همان راهِ جایگزین است: وقتی کاربر صریحاً گفته «بدون کد
    سفارش»، برداشتنِ کد از نام گیرنده یعنی کاری که او رد کرده — و خلاصه هم خلافش
    را می‌گوید.
    """
    rows: list[dict[str, Any]] = []
    for r in range(header_idx + 1, len(df)):
        get = lambda col, _r=r: _cell(df, _r, col)  # noqa: E731

        b_raw, c_raw, rn_raw = get(bc_col), get(code_col), get(row_col)
        b, c = _clean(b_raw), _clean(c_raw)

        # ستون «کد سفارش» جدا نبود یا خالی بود → کد را از داخل نام گیرنده بردار
        # (مثل «امیرحسین عاشوری ۳۰۶۱۷۶»)
        code_from_name = False
        if not c and name_for_code and name_col is not None:
            m = _RE_CODE_IN_TEXT.search(str(get(name_col) or "").translate(_FA2EN_TABLE))
            if m:
                c = m.group(1)
                code_from_name = True

        rn = _clean(rn_raw) or str(len(rows) + 1)
        has_digits_b = any(ch.isdigit() for ch in b)
        has_digits_c = any(ch.isdigit() for ch in c)
        if not has_digits_b and not has_digits_c:
            continue  # سطر خالی یا سطر «جمع کل» — رد می‌شود
        numeric_b = isinstance(b_raw, (int, float)) and not isinstance(b_raw, bool)
        rows.append(
            {
                "rownum": rn,
                "barcode": b,
                "code": c,
                "barcode_numeric": numeric_b,
                "barcode_raw": b_raw,
                "code_from_name": code_from_name,
                "sheet_row": r + 1,
            }
        )
    return rows


def _page_lines(page) -> list[list[tuple[float, float, str]]]:
    """Spans of one page, grouped into visual lines: ``(x0, x1, text)`` left→right.

    Grouping a line by its baseline is what makes this layout-independent; rounding y to
    0.1pt (the previous version) split one row into several «lines» whenever a cell used
    a different font or size, and those rows vanished from the output.
    """
    spans: list[tuple[float, float, float, str]] = []
    for block in page.get_text("rawdict").get("blocks", []):
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                text = "".join(ch["c"] for ch in span.get("chars", [])).strip().translate(_FA2EN_TABLE)
                if text:
                    x0, y0, x1 = (float(value) for value in span["bbox"][:3])
                    spans.append((y0, x0, x1, text))
    spans.sort(key=lambda item: (item[0], item[1]))
    lines: list[tuple[float, list[tuple[float, float, str]]]] = []
    for y0, x0, x1, text in spans:
        if lines and abs(lines[-1][0] - y0) <= _LINE_TOLERANCE:
            lines[-1][1].append((x0, x1, text))
        else:
            lines.append((y0, [(x0, x1, text)]))
    return [sorted(items) for _, items in lines]


def _pdf_barcode_index(line: list[tuple[float, float, str]]) -> int | None:
    """Which span of the line is the tracking barcode: the longest digit run.

    The previous version read fixed x bands (the printed form's own geometry). One
    changed margin, printer setting or form revision moved every cell, and the file came
    back as «ساختار PDF شناخته نشد» — while the 24-digit barcode was sitting right
    there. A length the shop calls a barcode wins over a longer run; then the longest
    run wins, so an order code or a phone number never becomes the barcode.
    """
    best_index, best_key = None, (0, 0)
    for index, (_x0, _x1, text) in enumerate(line):
        if not re.fullmatch(r"[0-9]+", text) or len(text) < _MIN_BARCODE_DIGITS:
            continue
        key = (int(len(text) in settings.barcode_lengths), len(text))
        if key > best_key:
            best_index, best_key = index, key
    return best_index


def _pdf_row(line: list[tuple[float, float, str]], index: int, page_number: int, count: int):
    """One printed table row → the same record shape the spreadsheet path produces."""
    _x0, _x1, barcode = line[index]
    number = ""
    number_index: int | None = None
    for position in range(len(line) - 1, index, -1):
        # ستون «ردیف» سمت راستِ بارکد است: عددِ کوتاهِ آخرِ سطر.
        if re.fullmatch(r"[0-9]{1,4}", line[position][2]):
            number, number_index = line[position][2], position
            break
    # کد سفارش در متنِ همان سطر است، هر جا که چاپ شده باشد (چیدمان‌ها فرق می‌کنند);
    # فقط خودِ بارکد و ستون ردیف کنار گذاشته می‌شوند.
    middle = " ".join(
        text for position, (_x0, _x1, text) in enumerate(line) if position not in (index, number_index)
    )
    mid_clean = _RE_DATE.sub("", middle)  # حذف تاریخ
    m = _RE_CODE_IN_TEXT.search(mid_clean)  # کد سفارش
    return {
        "rownum": number or str(count + 1),
        "barcode": barcode,
        "code": m.group(1) if m else "",
        "barcode_numeric": False,
        "barcode_raw": barcode,
        "code_from_name": False,
        "page": page_number,
        "sheet_row": None,
    }


def _read_pdf(path) -> list[dict[str, Any]]:
    """استخراج جدول از PDF خروجی سامانه — با هندسهٔ *خودِ* فایل، نه مختصات ثابت."""
    import pymupdf

    try:
        doc = pymupdf.open(path)
    except MemoryError as exc:
        raise _memory_error(exc) from exc
    except Exception as exc:
        logger.warning("pdf could not be opened: %s", path, exc_info=True)
        raise _unreadable("این فایل PDF سالم نیست — خراب یا نصفه دانلود شده", exc) from exc
    rows: list[dict[str, Any]] = []
    try:
        if len(doc) > 500:
            raise RowLimitError("بیش از ۵۰۰ صفحه در PDF؛ فایل را تقسیم کن")
        for page_number in range(len(doc)):
            for line in _page_lines(doc[page_number]):
                index = _pdf_barcode_index(line)
                if index is None:
                    continue
                _guard_rows(len(rows))
                rows.append(_pdf_row(line, index, page_number + 1, len(rows)))
    except MemoryError as exc:
        # «rawdict» یک صفحهٔ پرمتن می‌تواند صدها مگابایت شیء پایتون بسازد؛ پیش‌تر این
        # خطا بدون هیچ ترجمه‌ای به کاربر می‌رسید: «❌ خطا در پردازش: MemoryError».
        raise _memory_error(exc) from exc
    finally:
        doc.close()
    if not rows:
        raise UnreadableFileError(
            "در این PDF هیچ عدد بلندی (حداقل ۸ رقم) که بشود کد رهگیری باشد پیدا نشد. "
            "اگر PDF اسکن‌شده یا عکس است، خروجیِ متنی/جدولی همان گزارش را از سامانه بگیر."
        )
    return rows


# ---------------------------------------------------------------------------
# بررسی مشکلات
# ---------------------------------------------------------------------------
def _where(row: dict[str, Any], layout: Layout | None = None, col: int | None = None) -> str:
    """کجا این سطر را پیدا کنیم — لینک، نه حدس.

    The cell named is the one that is wrong (the barcode column, or the order-code
    column), so «برو به آن سلول» in Excel lands on the problem itself.
    """
    sheet = layout.sheet if layout is not None else ""
    sheet_row = row.get("sheet_row")
    if sheet and sheet_row:
        if col is None and layout is not None:
            col = layout.barcode_col
        return f"{sheet}!{_col_letter(col or 0)}{sheet_row}"
    if sheet_row:
        # A CSV has lines, not cells: saying «Sheet1!B2» would be a made-up address.
        return f"سطر {sheet_row} فایل"
    if row.get("page"):
        return f"صفحهٔ {row['page']}"
    return f"ردیف {row.get('rownum', '?')}"


def _barcode_state(row: dict[str, Any]) -> tuple[str, str]:
    """وضعیت بارکد یک سطر: ('ok'|'warn'|'error', شرح)

    هر سه حالت در گزارش ثبت می‌شود، اما فقط `ok` و `warn` وارد tracking.csv می‌شوند:
    اکسل اعداد ۲۴ رقمی را به float تبدیل و خراب می‌کند و آن عدد خراب در فایل
    خروجی کاملاً معقول به‌نظر می‌رسد — پس نوشتنش خطرناک‌تر از خالی‌گذاشتنش است.
    """
    b = row["barcode"]
    if not b:
        return "error", "بارکد خالی"
    raw = row.get("barcode_raw")
    if _destroyed(row):
        return "error", ("بارکد به‌صورت عدد ذخیره شده و دقتش از بین رفته — از فایل متنی/PDF اصلی استفاده کن")
    state, note = barcodes.classify(b)
    if state != "ok":
        return state, note
    if barcodes.number_is_suspicious(raw, b):
        return "warn", "بارکد در اکسل عددی است؛ اگر ۱۵ رقم را رد کند دقتش از بین می‌رود"
    return "ok", ""


def _order_code_state(row):
    c = row["code"]
    if not c:
        return "warn", "کد سفارش خالی (باید خودت تکمیل کنی)"
    if barcodes.order_code_is_valid(c):
        return "ok", ""
    if barcodes.order_code_is_short(c):
        return "warn", "کد سفارش ۵ رقمی (احتمالاً یک رقم جا افتاده)"
    return "error", f"کد سفارش نامعتبر: «{c}» (باید ۶ رقم باشد)"


def _repeat_label(rows: list[dict[str, Any]]) -> str:
    """«۲ و ۵ و ۸» — with a ceiling, because a barcode pasted down a whole column is
    a real export and the label is read by a human, not by a machine."""
    shown = [str(r["rownum"]) for r in rows[:3]]
    if len(rows) > 3:
        shown.append(f"+{len(rows) - 3} سطر دیگر")
    return " و ".join(shown)


def _analyze(rows: list[dict[str, Any]], layout: Layout | None = None) -> list[tuple]:
    """برمی‌گرداند: list of (ردیف نمایشی, شدت ❌/⚠️, شرح, بارکد, کد, لینک مبدأ)

    هر سطر هم برچسب می‌گیرد (`row['barcode_state']`, `row['code_state']`) تا
    `build_csv` بداند کدام داده قابل ارسال به سامانهٔ رهگیری است.
    """
    problems: list[tuple] = []
    #: duplicates keep the whole row, not just its label: the link has to point at
    #: both places the same barcode appears.
    bc_rows: dict[str, list[dict[str, Any]]] = {}
    code_rows: dict[str, list[dict[str, Any]]] = {}

    def cell(row: dict[str, Any], field_col: int | None) -> str:
        return _where(row, layout, field_col)

    for row in rows:
        n = row["rownum"]
        b, c = row["barcode"], row["code"]
        bc_state, bc_message = _barcode_state(row)
        row["barcode_state"] = bc_state
        row["barcode_note"] = bc_message
        if bc_message:
            problems.append(
                (
                    n,
                    "❌" if bc_state == "error" else "⚠️",
                    bc_message,
                    b,
                    c,
                    cell(row, layout.barcode_col if layout else None),
                )
            )
        if bc_state in ("ok", "warn"):
            bc_rows.setdefault(b, []).append(row)

        code_state, code_message = _order_code_state(row)
        row["code_state"] = code_state
        row["code_note"] = code_message
        if code_message:
            problems.append(
                (
                    n,
                    "❌" if code_state == "error" else "⚠️",
                    code_message,
                    b,
                    c,
                    cell(row, layout.code_col if layout else None),
                )
            )
        if code_state == "ok":
            code_rows.setdefault(c, []).append(row)

    # --- تکراری‌ها -----------------------------------------------------------
    # بارکد تکراری هشدار است نه خطا: یک بسته ممکن است دو سفارش داشته باشد و
    # کدش معتبر است. حذف هر دو سطر از CSV سفارش درست را از دست می‌داد.
    for b, found in bc_rows.items():
        if len(found) > 1:
            problems.append(
                (
                    _repeat_label(found),
                    "⚠️",
                    f"بارکد تکراری در {len(found)} سطر (در CSV می‌ماند؛ بررسی کن)",
                    b,
                    "",
                    " · ".join(cell(r, layout.barcode_col if layout else None) for r in found[:3]),
                )
            )
    for c, found in code_rows.items():
        if len(found) > 1:
            problems.append(
                (
                    _repeat_label(found),
                    "⚠️",
                    "کد سفارش تکراری",
                    "",
                    c,
                    " · ".join(cell(r, layout.code_col if layout else None) for r in found[:3]),
                )
            )

    return problems


# ---------------------------------------------------------------------------
# خروجی‌ها
# ---------------------------------------------------------------------------
def build_csv(rows) -> str:
    """فایل tracking.csv — فقط داده‌ای که معتبر است.

    سطرهای دارای خطا (بارکد خراب‌شده/ناصحیح) اینجا نمی‌آیند و به
    `needs-review.csv` می‌روند؛ کد سفارش خالی همچنان با سلول خالی می‌ماند تا
    خود کاربر تکمیلش کند.
    """
    out = ["order_id,tracking_code"]
    for row in rows:
        if row.get("barcode_state") not in ("ok", "warn"):
            continue
        code = row["code"] if row.get("code_state") in ("ok", "warn") else ""
        out.append(f"{code},{row['barcode']}")
    return "\n".join(out) + "\n"


def _safe_cell(value: object) -> str:
    text = str(value)
    return "'" + text if text.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")) else text


def _csv_join(cells: list[str]) -> str:
    """CSV فیلدِ ایمن: کاما/دوزاق/خط جدید در متنِ فارقی را نمی‌شکند."""
    out = []
    for cell in cells:
        text = _safe_cell(cell)
        if any(ch in text for ch in [",", '"', "\n", "\r"]):
            text = '"' + text.replace('"', '""') + '"'
        out.append(text)
    return ",".join(out)


def _destroyed(row: dict[str, Any]) -> bool:
    """بارکدی که اکسل عددش کرده و رقم‌هایش را خورده — یک‌جا، نه در سه جا."""
    raw = row.get("barcode_raw")
    return bool(row.get("barcode")) and barcodes.looks_float_destroyed(raw, row["barcode"])


def _fix_barcode(row: dict[str, Any]) -> str:
    """The barcode cell of the *editable* file: blank when the number is destroyed.

    Putting the destroyed digits back is how they get re-imported as if they were
    real, and the point of the fix file is that it is sent straight back to the bot.
    """
    return "" if _destroyed(row) else str(row.get("barcode") or "")


def _reason(row: dict[str, Any]) -> str:
    """Why this row is in the review file — the sentence from the analysis, not a label.

    «بارکد نامعتبر» would send the owner to problems.csv to learn anything; the row
    already knows what was wrong with it.
    """
    parts = []
    if _destroyed(row):
        # The analysis sentence says «use the text/PDF export next time», which is advice
        # about the *next* file; the cell needs what to do with this one.
        parts.append("بارکد در اکسل خراب شده — خودت از فایل اصلی بردار")
    elif row.get("barcode_state") in ("error", "warn"):
        parts.append(str(row.get("barcode_note") or "بارکد نامعتبر"))
    if row.get("code_state") in ("error", "warn"):
        parts.append(str(row.get("code_note") or "کد سفارش نامعتبر"))
    return " + ".join(parts)


def review_rows(rows) -> list[dict[str, Any]]:
    """سطرهایی که باید دیده شوند: خطاها، و آن هشدارهایی که در CSV هم مانده‌اند.

    نام فایل «needs-review» است نه «needs-fix»، چون یک EAN-13 که پذیرفتیم اینجا می‌آید
    *به‌خاطر اینکه* ارزش یک نگاه دوباره را دارد؛ حذفش از فهرست یعنی پنهان‌کردنِ همان
    حالتی که کدِ معتبرْ کدِ رهگیری نیست.
    """
    return [r for r in rows if _reason(r)]


def unwritten_rows(rows) -> list[dict[str, Any]]:
    """سطرهایی که اصلاً در tracking.csv ننوشته شدند (بارکدشان خطادار است)."""
    return [r for r in rows if r.get("barcode_state") == "error"]


def build_review_csv(rows, layout: Layout | None = None) -> str | None:
    """سطرهای نیازمند اصلاح، با دلیل و لینک سطر — برای اینکه گم نشوند."""
    lines = ["row,source_cell,order_id,barcode,reason"]
    any_row = False
    for row in review_rows(rows):
        any_row = True
        lines.append(
            _csv_join(
                [
                    str(row["rownum"]),
                    _where(row, layout),
                    row["code"],
                    row["barcode"],
                    _reason(row),
                ]
            )
        )
    if not any_row:
        return None
    return "\n".join(lines) + "\n"


def build_review_workbook(rows, layout: Layout | None = None) -> bytes | None:
    """همان سطرها در اکسل — با ستون بارکدِ «متن»، تا اصلاح‌کردن خرابشان نکند.

    سرستون‌ها عمداً نام‌های خودِ جدول هستند («ردیف/بارکد/کد سفارش»): کاربر در اکسل
    بارکد را اصلاح می‌کند و **همان فایل** را دوباره برای ربات می‌فرستد؛ مسیرِ
    خواندنِ دیگری لازم نیست.
    """
    items = review_rows(rows)
    if not items:
        return None
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "needs-review"
    headers = ["ردیف", "بارکد", "کد سفارش", "دلیل", "محل در فایل اصلی"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in items:
        ws.append(
            [
                str(row["rownum"]),
                _fix_barcode(row),
                str(row["code"]),
                _reason(row),
                _where(row, layout),
            ]
        )
    for r in range(2, len(items) + 2):
        for cell in ws[r]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
        for col in (1, 2, 3):
            ws.cell(row=r, column=col).number_format = "@"  # متن، نه عدد
            ws.cell(row=r, column=col).alignment = Alignment(horizontal="left")
    for letter, width in zip("ABCDE", (8, 30, 12, 22, 20), strict=True):
        ws.column_dimensions[letter].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_problems_csv(problems) -> str | None:
    """`problems.csv`: هر مشکل با لینک سطرِ مبدأ، نه یک فایل متنی بی‌نقشه."""
    if not problems:
        return None
    lines = ["severity,source_row,source_cell,reason,order_id,barcode"]
    for n, sev, desc, b, c, where in problems:
        lines.append(_csv_join([sev, str(n), where, desc, c, b]))
    return "\n".join(lines) + "\n"


def build_summary(rows, problems, fname, layout: Layout | None = None, notes: tuple[str, ...] = ()) -> str:
    n = len(rows)
    usable = sum(1 for r in rows if r.get("barcode_state") in ("ok", "warn"))
    dropped = len(unwritten_rows(rows))
    c_valid = sum(1 for r in rows if barcodes.order_code_is_valid(r["code"]))
    errs = [p for p in problems if p[1] == "❌"]
    warns = [p for p in problems if p[1] == "⚠️"]

    lines = [
        f"📄 فایل: {fname}",
        f"🔢 تعداد ردیف: {n}",
        f"✅ ردیف قابل استفاده در tracking.csv: {usable} از {n}",
        f"📦 کد سفارش معتبر: {c_valid} از {n}",
        f"❌ خطا: {len(errs)}   ⚠️ هشدار: {len(warns)}",
    ]
    schema = layout.schema_line() if layout is not None else ""
    if schema:
        lines += ["", f"🗺️ ستون‌ها: {schema}"]
    if notes:
        lines += ["", *[f"ℹ️ {note}" for note in notes]]
    if dropped:
        lines += [
            "",
            f"🚫 {dropped} ردیف به‌دلیل بارکد نامعتبر در tracking.csv ننوشته شد"
            " (در needs-review.csv با دلیل آمده؛ برای اصلاح در اکسل needs-review.xlsx را باز کن"
            " و بعد از اصلاح، همان فایل را دوباره بفرست).",
        ]
    if problems:
        lines += ["", "مشکلات (نمونه):"]
        for entry in problems[:8]:
            rn, sev, desc = entry[0], entry[1], entry[2]
            lines.append(f"{sev} ردیف {rn}: {desc}")
        if len(problems) > 8:
            lines.append(f"… و {len(problems) - 8} مورد دیگر (در فایل problems.csv)")
        lines += ["", "📎 فایل tracking.csv ضمیمه شد — کدهای خالی را خودت تکمیل کن."]
    else:
        lines += ["", "🎉 هیچ مشکلی پیدا نشد! فایل tracking.csv آماده است."]
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# نتیجه
# ---------------------------------------------------------------------------
@dataclass(frozen=True)
class Report:
    """Everything one file produced. ``questions`` non-empty means: ask, don't write."""

    csv_text: str = ""
    summary: str = ""
    problems_csv: str | None = None
    review_csv: str | None = None
    review_xlsx: bytes | None = None
    #: rows the owner must look at (errors *and* warnings)
    needs_review: int = 0
    #: rows that could not be written to tracking.csv at all
    dropped: int = 0
    layout: Layout | None = None
    questions: tuple[Question, ...] = ()
    fname: str = ""
    notes: tuple[str, ...] = ()
    rows: int = 0
    usable: int = 0
    errors: int = 0
    warnings: int = 0

    @property
    def needs_answer(self) -> bool:
        return bool(self.questions)

    def as_dict(self) -> dict[str, Any]:
        """What the ledger may keep — the files themselves are never stored."""
        return {
            "rows": self.rows,
            "usable": self.usable,
            "errors": self.errors,
            "warnings": self.warnings,
            "needs_review": self.needs_review,
            "dropped": self.dropped,
            "file": self.fname,
        }


# ---------------------------------------------------------------------------
# نقطه‌ی ورود اصلی
# ---------------------------------------------------------------------------
def _guard_rows(count: int) -> None:
    """رد کردنِ فایلِ بزرگ، با عددی که راست است.

    Reading stops at ``MAX_ROWS + 1`` rows on purpose (a 25 MB xlsx can expand to
    gigabytes), so the honest sentence is «بیش از N ردیف» — the previous message printed
    the truncated count («حداقل ۴») as if it were the file's real size.
    """
    limit = int(settings.max_rows)
    if limit > 0 and count > limit:
        raise RowLimitError(
            f"فایل بیش از {limit:,} ردیف دارد (خواندن روی سقف متوقف شد) و سقف {limit:,} ردیف "
            "(MAX_ROWS) را رد کرده است. فایل را به دو یا چند بخش تقسیم کن؛ "
            "پردازشِ این حجم، ربات را برای همه slow می‌کند."
        )


def process_file(path, fname=None, *, layout: Layout | dict[str, Any] | None = None) -> Report:
    """Read one order file → ``Report``.

    ``layout`` is what the user already chose (see :func:`scan`); without it the
    header row and the columns are detected, and a file whose headers are
    ambiguous comes back with ``questions`` instead of a guessed answer.
    """
    ext = os.path.splitext(str(path))[1].lower()
    fname = fname or os.path.basename(str(path))
    chosen = layout if isinstance(layout, Layout) else (Layout.from_dict(layout) if layout else None)

    if ext == ".pdf":
        rows = _read_pdf(path)
        _guard_rows(len(rows))
        reader_notes: tuple[str, ...] = ()
        report_layout = chosen or Layout(header_row=-1)
        questions: tuple[Question, ...] = ()
    elif ext in (".xlsx", ".xlsm", ".csv"):
        df, sheet, reader_notes = _read_table(path, ext)
        _guard_rows(len(df))
        answers: dict[str, int | None] = {}
        if chosen is not None:
            # A column of ``None`` means two different things: never looked for, or
            # switched off on purpose. ``decided`` is what tells them apart, and only
            # the second one must not be re-detected (re-detecting it asks again).
            answers = {"header_row": chosen.header_row}
            for field, column in ((FIELD_BARCODE, chosen.barcode_col), (FIELD_CODE, chosen.code_col)):
                if column is not None:
                    answers[field] = column
                elif field in dict(chosen.decided):
                    answers[field] = -1
        report_layout, questions = scan_table(
            df, answers, sheet=sheet, prior_labels=dict(chosen.decided) if chosen else None
        )
        if questions:
            # Ask before writing anything: a wrong barcode column produces an
            # empty file, a wrong order-code column produces a *wrong* one.
            return Report(questions=questions, layout=report_layout, fname=fname)
        if report_layout.barcode_col is None:
            raise ValueError(
                "ستون «بارکد» در فایل پیدا نشد. فایل باید جدول سفارش‌ها "
                "(ردیف / بارکد / تاریخ ثبت / نام گیرنده / کد سفارش / …) باشد."
            )
        # «کد سفارش» سه حالت دارد: ستونش هست، نیست ولی از نام گیرنده درمی‌آید، یا
        # کاربر خودش گفته ندارد. سومین با دومی فرق دارد و همان را از راهِ نام می‌بندد.
        skipped_code = report_layout.code_col is None and FIELD_CODE in dict(report_layout.decided)
        rows = _collect_rows(
            df,
            report_layout.header_row,
            report_layout.barcode_col,
            report_layout.code_col,
            report_layout.row_col,
            report_layout.name_col,
            name_for_code=not skipped_code,
        )
    else:
        raise ValueError(f"فرمت «{ext}» پشتیبانی نمی‌شود (فقط xlsx / xlsm / csv / pdf).")

    if not rows:
        raise ValueError("هیچ ردیف داده‌ای در فایل پیدا نشد.")

    # یادداشت‌های خودِ خواندن (کدگذاری/جداکننده/برگه) اول می‌آیند: کاربر باید بداند
    # فایلش چطور خوانده شد، نه فقط اینکه چه چیزی داخلش بود.
    notes: list[str] = list(reader_notes)
    if report_layout.barcode_col is not None and report_layout.code_col is None:
        lifted = sum(1 for r in rows if r.get("code_from_name"))
        if skipped_code:
            notes.append(
                "ستون کد سفارش را خودت «بدون» انتخاب کردی؛ هیچ کدی (حتی از نام گیرنده) خوانده نشد."
            )
        elif lifted:
            notes.append(
                f"{lifted} کد سفارش از ستونِ نام گیرنده برداشته شد — این‌ها را چشم‌بسته تأیید نکن."
            )
        else:
            notes.append("ستون «کد سفارش» نبود و در نام گیرنده هم کدی پیدا نشد؛ کدِ همهٔ ردیف‌ها خالی است.")
    destroyed = sum(1 for r in rows if _destroyed(r))
    if destroyed:
        notes.append(
            f"{destroyed} بارکد در اکسل به‌صورت عدد ذخیره شده بود و رقم‌هایش رفته است — "
            "خروجیِ متنی (CSV) یا همان PDFِ سامانه را بفرست، اکسل نه."
        )

    problems = _analyze(rows, report_layout)
    usable = sum(1 for r in rows if r.get("barcode_state") in ("ok", "warn"))
    return Report(
        csv_text=build_csv(rows),
        summary=build_summary(rows, problems, fname, report_layout, tuple(notes)),
        problems_csv=build_problems_csv(problems),
        review_csv=build_review_csv(rows, report_layout),
        review_xlsx=build_review_workbook(rows, report_layout),
        layout=report_layout,
        notes=tuple(notes),
        fname=fname,
        rows=len(rows),
        usable=usable,
        errors=sum(1 for p in problems if p[1] == "❌"),
        warnings=sum(1 for p in problems if p[1] == "⚠️"),
        needs_review=len(review_rows(rows)),
        dropped=len(unwritten_rows(rows)),
    )


__all__ = [
    "FIELD_BARCODE",
    "FIELD_CODE",
    "Column",
    "Layout",
    "MemoryLimitError",
    "Question",
    "Report",
    "RowLimitError",
    "UnreadableFileError",
    "build_csv",
    "build_problems_csv",
    "build_review_csv",
    "build_review_workbook",
    "build_summary",
    "process_file",
    "review_rows",
    "scan_table",
    "unwritten_rows",
]
